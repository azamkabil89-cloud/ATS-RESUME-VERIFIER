import io
import json
import re
from datetime import datetime
from xml.sax.saxutils import escape

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

MODEL_NAME = "gemini-2.5-flash"
ALLOWED_TYPES = ["pdf", "docx", "txt"]
MAX_FILE_MB = 5
MIN_TEXT_CHARS = 100

st.set_page_config(page_title="ATS Resume Checker", page_icon="📄", layout="centered")


class ResumeError(Exception):
    """A user-friendly problem with the uploaded resume."""


# ---------------- API key ----------------
def get_api_key():
    try:
        return st.secrets.get("GEMINI_API_KEY")
    except Exception:
        return None


def validate_api_key(key):
    if not key or not key.strip():
        return "No Gemini API key found. Paste your key above, or add GEMINI_API_KEY to Streamlit Secrets. Get a free key at https://aistudio.google.com/apikey"
    if len(key.strip()) < 20 or " " in key.strip():
        return "That API key doesn't look right. Copy it again from https://aistudio.google.com/apikey without extra spaces."
    return None


# ---------------- File reading ----------------
def extract_text(uploaded_file):
    name = uploaded_file.name
    ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""

    if ext not in ALLOWED_TYPES:
        raise ResumeError(f"'{name}' is not supported. Please upload a PDF, DOCX or TXT file. (Old .doc files: open in Word and 'Save As' .docx.)")

    data = uploaded_file.getvalue()
    if len(data) == 0:
        raise ResumeError(f"'{name}' is empty (0 bytes). Please upload a file that contains your resume.")
    if len(data) > MAX_FILE_MB * 1024 * 1024:
        raise ResumeError(f"'{name}' is larger than {MAX_FILE_MB} MB. Please upload a smaller file.")

    try:
        if ext == "pdf":
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                try:
                    reader.decrypt("")
                except Exception:
                    pass
                if reader.is_encrypted and not _can_read_pdf(reader):
                    raise ResumeError("This PDF is password-protected. Please remove the password and upload it again.")
            if len(reader.pages) == 0:
                raise ResumeError("This PDF has no pages.")
            text = "\n".join((page.extract_text() or "") for page in reader.pages)
        elif ext == "docx":
            doc = Document(io.BytesIO(data))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            text = "\n".join(parts)
        else:
            text = _decode_txt(data)
    except ResumeError:
        raise
    except Exception:
        raise ResumeError(f"'{name}' could not be read. It may be damaged or not a real {ext.upper()} file. Try re-saving or exporting it again.")

    text = text.strip()
    if not text:
        raise ResumeError("No text was found in this file. If it is a scanned image or photo, please upload a text-based PDF or a DOCX file.")
    if len(text) < MIN_TEXT_CHARS:
        raise ResumeError(f"Only {len(text)} characters of text were found. This is too short to be a resume. If it's a scanned PDF, upload a text-based version.")
    return text


def _can_read_pdf(reader):
    try:
        _ = reader.pages[0].extract_text()
        return True
    except Exception:
        return False


def _decode_txt(data):
    for enc in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ResumeError("This text file uses an unknown encoding. Please save it as UTF-8 and try again.")


# ---------------- AI analysis ----------------
def build_prompt(resume_text, job_description):
    jd_part = (
        f"JOB DESCRIPTION:\n{job_description}\n"
        if job_description.strip()
        else "No job description was given. Judge the resume against general ATS best practices.\n"
    )
    return f"""You are an expert ATS (Applicant Tracking System) and recruiter.
Analyse the resume below and reply ONLY with valid JSON in exactly this shape:
{{
  "ats_score": <integer 0-100>,
  "summary": "<2-3 sentence overall verdict>",
  "strengths": ["..."],
  "weaknesses": ["..."],
  "missing_keywords": ["..."],
  "improvements": ["<specific, actionable suggestion>"],
  "section_scores": {{
    "formatting": <0-100>,
    "keywords": <0-100>,
    "experience": <0-100>,
    "skills": <0-100>,
    "education": <0-100>
  }}
}}

{jd_part}
RESUME:
{resume_text}
"""


def parse_json(text):
    text = (text or "").strip()
    text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text)
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", text, re.DOTALL)
        if match:
            return json.loads(match.group(0))
        raise


def friendly_api_error(e):
    msg = str(e).lower()
    if "api key" in msg or "api_key_invalid" in msg or "permission" in msg or "401" in msg or "403" in msg:
        return "Your Gemini API key was rejected. Check that it is correct and active at https://aistudio.google.com/apikey"
    if "quota" in msg or "429" in msg or "resource_exhausted" in msg:
        return "Gemini usage limit reached. Wait a minute and try again, or check your quota in Google AI Studio."
    if "not found" in msg or "404" in msg:
        return f"The model '{MODEL_NAME}' is not available for your key."
    if "timeout" in msg or "deadline" in msg or "connection" in msg or "503" in msg:
        return "Couldn't reach Gemini right now. Check your internet connection and try again."
    return f"Analysis failed: {e}"


def analyse(api_key, resume_text, job_description):
    client = genai.Client(api_key=api_key.strip())
    response = client.models.generate_content(
        model=MODEL_NAME,
        contents=build_prompt(resume_text, job_description),
        config=types.GenerateContentConfig(temperature=0.2, response_mime_type="application/json"),
    )
    result = parse_json(response.text)
    if not isinstance(result, dict) or "ats_score" not in result:
        raise ValueError("The AI returned an unexpected answer. Please try again.")
    try:
        result["ats_score"] = max(0, min(int(result["ats_score"]), 100))
    except (TypeError, ValueError):
        result["ats_score"] = 0
    return result


# ---------------- PDF report ----------------
def _clean(text):
    """Make text safe for ReportLab's built-in fonts (Latin-1) and XML markup."""
    replacements = {"\u2018": "'", "\u2019": "'", "\u201c": '"', "\u201d": '"',
                    "\u2013": "-", "\u2014": "-", "\u2022": "-", "\u2026": "...", "\u00a0": " "}
    text = str(text)
    for k, v in replacements.items():
        text = text.replace(k, v)
    text = text.encode("latin-1", errors="ignore").decode("latin-1")
    return escape(text)


def build_pdf_report(result, file_name):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm,
                            topMargin=18 * mm, bottomMargin=18 * mm, title="ATS Resume Report")
    styles = getSampleStyleSheet()
    title = ParagraphStyle("T", parent=styles["Title"], textColor=colors.HexColor("#1F3A5F"))
    h2 = ParagraphStyle("H", parent=styles["Heading2"], textColor=colors.HexColor("#1F3A5F"), spaceBefore=10)
    body = ParagraphStyle("B", parent=styles["BodyText"], fontSize=10.5, leading=15)
    small = ParagraphStyle("S", parent=body, fontSize=9, textColor=colors.grey)

    score = result.get("ats_score", 0)
    score_color = "#2E7D32" if score >= 80 else "#ED6C02" if score >= 60 else "#C62828"

    flow = [
        Paragraph("ATS Resume Analysis Report", title),
        Paragraph(f"File: {_clean(file_name)} &nbsp;|&nbsp; Date: {datetime.now():%d %b %Y, %H:%M}", small),
        Spacer(1, 10),
        Paragraph(f'<font size="28" color="{score_color}"><b>{score}</b></font><font size="14"> / 100 ATS Score</font>', body),
        Spacer(1, 14),
        Paragraph(_clean(result.get("summary", "")), body),
    ]

    sections = result.get("section_scores") or {}
    if sections:
        flow.append(Paragraph("Section Scores", h2))
        rows = [["Section", "Score"]] + [[_clean(k).title(), str(v)] for k, v in sections.items()]
        table = Table(rows, colWidths=[110 * mm, 40 * mm], hAlign="LEFT")
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3A5F")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F5F9")]),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#C9D3DF")),
            ("ALIGN", (1, 0), (1, -1), "CENTER"),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        flow.append(table)

    for heading, key in [("Strengths", "strengths"), ("Weaknesses", "weaknesses"),
                         ("Missing Keywords", "missing_keywords"), ("Improvements", "improvements")]:
        items = result.get(key) or []
        flow.append(Paragraph(heading, h2))
        if not items:
            flow.append(Paragraph("None", body))
        for item in items:
            flow.append(Paragraph(f"&bull; {_clean(item)}", body))

    doc.build(flow)
    return buffer.getvalue()


# ---------------- UI ----------------
def show_list(title, items):
    st.subheader(title)
    if not items:
        st.write("None")
    for item in items:
        st.markdown(f"- {item}")


st.title("📄 ATS Resume Checker")
st.write("Upload your resume to get an ATS score, suggestions, and a downloadable PDF report.")

api_key = get_api_key()
if not api_key:
    api_key = st.text_input("Enter your Gemini API key", type="password",
                            help="Get a free key at https://aistudio.google.com/apikey")

uploaded = st.file_uploader(f"Upload resume (PDF, DOCX or TXT, max {MAX_FILE_MB} MB)", type=ALLOWED_TYPES)
job_description = st.text_area("Paste job description (optional, gives better results)", height=150)

if st.button("Analyse Resume", type="primary"):
    st.session_state.pop("result", None)
    key_error = validate_api_key(api_key)
    if key_error:
        st.error(key_error)
    elif uploaded is None:
        st.error("Please upload your resume first (PDF, DOCX or TXT).")
    else:
        try:
            resume_text = extract_text(uploaded)
        except ResumeError as e:
            st.error(str(e))
        else:
            with st.spinner("Analysing your resume..."):
                try:
                    st.session_state["result"] = analyse(api_key, resume_text, job_description)
                    st.session_state["file_name"] = uploaded.name
                except json.JSONDecodeError:
                    st.error("The AI returned an unreadable answer. Please click Analyse again.")
                except Exception as e:
                    st.error(friendly_api_error(e))

result = st.session_state.get("result")
if result:
    score = result["ats_score"]
    st.metric("ATS Score", f"{score} / 100")
    st.progress(score / 100)
    if score >= 80:
        st.success("Great! Your resume is ATS-friendly.")
    elif score >= 60:
        st.warning("Good, but there is room for improvement.")
    else:
        st.error("Your resume needs significant improvement.")

    st.write(result.get("summary", ""))

    sections = result.get("section_scores") or {}
    if sections:
        st.subheader("Section Scores")
        cols = st.columns(len(sections))
        for col, (name, value) in zip(cols, sections.items()):
            col.metric(name.title(), value)

    show_list("✅ Strengths", result.get("strengths", []))
    show_list("⚠️ Weaknesses", result.get("weaknesses", []))
    show_list("🔑 Missing Keywords", result.get("missing_keywords", []))
    show_list("🛠️ Improvements", result.get("improvements", []))

    try:
        file_name = st.session_state.get("file_name", "resume")
        pdf_bytes = build_pdf_report(result, file_name)
        st.download_button("⬇️ Download PDF Report", data=pdf_bytes,
                           file_name=f"ATS_Report_{file_name.rsplit('.', 1)[0]}.pdf",
                           mime="application/pdf", type="primary")
    except Exception as e:
        st.error(f"Could not create the PDF report: {e}")
